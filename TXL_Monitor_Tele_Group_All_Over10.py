# ==========================================================
# KIEM_SOAT_PHIEU_TXL_TELE_GROUP.py
# ==========================================================
import os
import re
import html
import requests
import pandas as pd

from datetime import datetime, timedelta

import tkinter as tk
from tkinter import ttk
from tkinter import filedialog
from tkinter import messagebox

from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    Alignment,
    PatternFill,
    Border,
    Side
)

from openpyxl.utils import get_column_letter


# ==========================================================
# APP INFO
# ==========================================================

APP_TITLE = "KIỂM SOÁT PHIẾU TIỀN XỬ LÝ"


# ==========================================================
# FILE PATTERN
# ==========================================================

BH_PATTERN = r"Bao_hong_ton_\d{12}\.(xlsx|xls)$"

HT_PATTERN = r"Ton_Khieu_Nai_\d{12}\.(xlsx|xls)$"


# ==========================================================
# SERVICE KEYWORDS
# ==========================================================

SERVICE_KEYWORDS = [
    "colocation", "email", "hosting", "ereceipt", "pharmacy", "clinic", "hkd", "i-market", "eheza", "hóa đơn điện tử", "ceca", "ca-tax",
    "vnedu", "invoice", "vnpt ca", "bhxh", "bmis", "econtract", "smartca", "tsa", "accounting", "check", "mobile id", "máy tính tiền",
    "smartads", "asme", "ezozo", "onebusiness", "ialert", "eticket", "wasspro", "salespro", "meeting", "ssl", "wifi (hợp tác)", "tên miền",
    "domain", "cdn", "sd-wan", "vcc", "sllđt", "sms điều hành", "mooc", "tracking", "posio", "holio", "platform", "smartedu", "e-learning",
]

# ==========================================================
# DONVI KEYWORDS
# ==========================================================

DONVI_KEYWORDS = [
    "mb", "miền bắc", "mien bac",
    "mn", "miền nam", "mien nam",
    "mt", "miền trung", "mien trung",
    "cntt&dvs",
]

# ==========================================================
# MAP DON VI
# ==========================================================

MAP_DONVI = [
    (["mb", "miền bắc", "mien bac"], "P.HTKH MB"),
    (["mn", "miền nam", "mien nam"], "P.HTKH MN"),
    (["mt", "miền trung", "mien trung"], "P.HTKH MT"),
]


# ==========================================================
# DISPLAY COLUMNS BÁO HỎNG
# ==========================================================

DISPLAY_COLUMNS_BH = [

    ("Tỉnh/Thành phố", "tentinh"),
    ("Khách hàng", "ten_tb"),
    ("Mã thuê bao", "ma_tb"),
    ("Dịch vụ", "loaihinh_tb"),
    ("Đơn vị mở phiếu", "ten_dv"),
    ("Người mở phiếu", "ten_nv"),
    ("Mã báo hỏng", "ma_bh"),
    ("Nội dung báo hỏng", "ghichu_hong"),
    ("Liên hệ khách hàng", "dienthoai_lh"),
    ("Trạng thái phiếu", "ten_trangthai"),
    ("Trạng thái xử lý", "trangthai_bh"),
    ("Đơn vị xử lý", "ten_dv_xl"),
    ("Bộ phận đang xử lý", "ten_dv_dang_th"),
    ("Thời gian bắt đầu", "ngay_bh"),
    ("Thời gian đã xử lý", "thoigian_xl_bh"),
    ("SLA_TXL", "sla_bh"),

]

# ==========================================================
# DISPLAY COLUMNS HỖ TRỢ
# ==========================================================

DISPLAY_COLUMNS_HT = [

    ("Tỉnh/Thành phố", "TENTINH"),
    ("Mã khách hàng", "MA_KH"),
    ("Mã thuê bao", "MA_TB"),
    ("Dịch vụ", "LOAIHINH_TB"),
    ("Đơn vị mở phiếu", "TEN_DV"),
    ("Người mở phiếu", "TEN_NV"),
    ("Mã hỗ trợ", "MA_HT"),
    ("Nội dung hỗ trợ", "YC_HOTRO"),
    ("Liên hệ khách hàng", "DIENTHOAI_LH"),
    ("Trạng thái phiếu", "TRANGTHAI_HT"),
    ("Ngày tiếp nhận yêu cầu", "NGAY_TN"),
    ("Đơn vị xử lý", "DONVI_XULY"),
    ("Thời gian đã xử lý", "thoigian_xl_ht"),
    ("SLA_TXL", "sla_ht"),

]

# ==========================================================
# EXCEL STYLE
# ==========================================================

HEADER_FILL = PatternFill(
    "solid",
    "FFD966"
)

TOTAL_FILL = PatternFill(
    "solid",
    "D9D9D9"
)

YELLOW_FILL = PatternFill(
    "solid",
    "FFFF00"
)

RED_FILL = PatternFill(
    "solid",
    "F4CCCC"
)

THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin")
)


# ==========================================================
# TELEGRAM
# ==========================================================

# Do not commit Telegram credentials. Configure them in the environment or
# through the application settings before sending messages.
BOT_TOKEN = os.getenv("TXL_TELEGRAM_BOT_TOKEN", "")
GROUP_CHAT_ID = os.getenv("TXL_TELEGRAM_GROUP_CHAT_ID", "")

FILE_DATA_TIME = None

# ==========================================================
# HELPER FUNCTIONS
# ==========================================================
def normalize_name(name):

    return (
        str(name)
        .strip()
        .upper()
    )
    
def is_valid_bh_file(filepath):
    """
    Bao_hong_ton_123456789012.xlsx
    """
    filename = os.path.basename(filepath)

    return (
        re.fullmatch(
            BH_PATTERN,
            filename,
            flags=re.IGNORECASE
        )
        is not None
    )


# ==========================================================
# VALIDATE HT FILE
# ==========================================================

def is_valid_ht_file(filepath):
    """
    Ton_Khieu_Nai_123456789012.xlsx
    """

    filename = os.path.basename(filepath)

    return (
        re.fullmatch(
            HT_PATTERN,
            filename,
            flags=re.IGNORECASE
        )
        is not None
    )

# ==========================================================
# GET FILE DATETIME
# ==========================================================

def get_file_datetime(filepath):

    filename = os.path.basename(filepath)

    match = re.search(
        r"(\d{12})",
        filename
    )

    if not match:
        return datetime.now()

    s = match.group(1)

    try:

        return datetime.strptime(
            s,
            "%Y%m%d%H%M"
        )

    except:

        return datetime.now()

# ==========================================================
# MAP DON VI
# ==========================================================

def map_donvi(text):

    text = str(text).lower()

    for keywords, donvi in MAP_DONVI:

        for keyword in keywords:

            if keyword in text:
                return donvi

    return "KHAC"


# ==========================================================
# CHECK DON VI KEYWORD
# ==========================================================

def has_donvi_keyword(text):

    text = str(text).lower()

    return any(
        keyword in text
        for keyword in DONVI_KEYWORDS
    )


# ==========================================================
# CHECK SERVICE --> CHUYỂN SANG GIÁM SÁT ALL DỊCH VỤ
# ==========================================================

#def is_service_match(service_name):

#    service_name = str(service_name).lower()

#    return any(
#        keyword in service_name
#        for keyword in SERVICE_KEYWORDS
#    )


# ==========================================================
# PARSE DATETIME
# ==========================================================

def parse_datetime(value):

    if pd.isna(value):
        return pd.NaT

    try:
        return pd.to_datetime(
            value,
            dayfirst=True,
            errors="coerce"
        )
    except:
        return pd.NaT


# ==========================================================
# SLA
# ==========================================================

def calculate_sla_bh(minutes):

    try:
        minutes = float(minutes)
    except:
        return ""

    if minutes < 10:
        return "TXL < 10"

    elif minutes <= 20:
        return "10 ≤ TXL ≤ 20"

    elif minutes <= 40:
        return "20 < TXL ≤ 40"

    else:
        return "Đã quá hạn"


# ==========================================================
# FORMAT DATETIME
# ==========================================================

def format_datetime(dt):

    if pd.isna(dt):
        return ""

    try:
        return dt.strftime(
            "%d/%m/%Y %H:%M:%S"
        )
    except:
        return ""


# ==========================================================
# AUTO WIDTH
# ==========================================================

def autosize_columns(ws):

    header_row = 3

    for col in range(
        1,
        ws.max_column + 1
    ):

        header = str(
            ws.cell(
                header_row,
                col
            ).value or ""
        )

        width = max(
            len(header) + 4,
            15
        )

        ws.column_dimensions[
            get_column_letter(col)
        ].width = width

# ==========================================================
# EXCEL HEADER STYLE
# ==========================================================

def style_header_row(ws, row_number):

    for cell in ws[row_number]:

        cell.fill = HEADER_FILL

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        cell.border = THIN_BORDER


# ==========================================================
# APPLY ROW COLOR
# ==========================================================

def apply_sla_fill(ws, row_number, sla_value):

    if sla_value == "20 < TXL ≤ 40":

        fill = YELLOW_FILL

    elif sla_value == "Đã quá hạn":

        fill = RED_FILL

    else:

        return

    for cell in ws[row_number]:

        cell.fill = fill


# ==========================================================
# GET EXPORT TIMESTAMP
# ==========================================================

def get_export_timestamp():

    return datetime.now().strftime(
        "%H:%M ngày %d/%m/%Y"
    )


# ==========================================================
# BUILD SUMMARY TABLE
# ==========================================================

def build_summary_dataframe(df):

    rows = []

    for donvi in [

        "P.HTKH MB",
        "P.HTKH MT",
        "P.HTKH MN"

    ]:

        g = df[
            df["DONVI"] == donvi
        ]
        
        cnt_10 = (
            g["sla_bh"]
            .eq("TXL < 10")
            .sum()
        )

        cnt_20 = (
            g["sla_bh"]
            .eq("10 ≤ TXL ≤ 20")
            .sum()
        )

        cnt_40 = (
            g["sla_bh"]
            .eq("20 < TXL ≤ 40")
            .sum()
        )

        cnt_qh = (
            g["sla_bh"]
            .eq("Đã quá hạn")
            .sum()
        )

        total = len(g)

        rows.append({

            "Đơn vị": donvi,
            
            "TXL < 10": cnt_10,

            "10 ≤ TXL ≤ 20": cnt_20,

            "20 < TXL ≤ 40": cnt_40,

            "Đã quá hạn": cnt_qh,

            "Tổng": total

        })

    summary = pd.DataFrame(rows)

    total_row = {

        "Đơn vị": "TỔNG",
        
        "TXL < 10":
            summary["TXL < 10"].sum(),

        "10 ≤ TXL ≤ 20":
            summary["10 ≤ TXL ≤ 20"].sum(),

        "20 < TXL ≤ 40":
            summary["20 < TXL ≤ 40"].sum(),

        "Đã quá hạn":
            summary["Đã quá hạn"].sum(),

        "Tổng":
            summary["Tổng"].sum()

    }

    summary = pd.concat(
        [
            summary,
            pd.DataFrame([total_row])
        ],
        ignore_index=True
    )

    return summary
    
# ==========================================================
# BUILD SUMMARY HT
# ==========================================================

def build_summary_dataframe_ht(df):

    rows = []

    for donvi in [

        "P.HTKH MB",
        "P.HTKH MT",
        "P.HTKH MN"

    ]:

        g = df[
            df["DONVI"] == donvi
        ]
        
        cnt_10 = (
            g["sla_ht"]
            .eq("TXL < 10")
            .sum()
        )

        cnt_20 = (
            g["sla_ht"]
            .eq("10 ≤ TXL ≤ 20")
            .sum()
        )

        cnt_40 = (
            g["sla_ht"]
            .eq("20 < TXL ≤ 40")
            .sum()
        )

        cnt_qh = (
            g["sla_ht"]
            .eq("Đã quá hạn")
            .sum()
        )

        total = len(g)

        rows.append({

            "Đơn vị": donvi,
            
            "TXL < 10": cnt_10,

            "10 ≤ TXL ≤ 20": cnt_20,

            "20 < TXL ≤ 40": cnt_40,

            "Đã quá hạn": cnt_qh,

            "Tổng": total

        })

    summary = pd.DataFrame(rows)

    total_row = {

        "Đơn vị": "TỔNG",
        
        "TXL < 10":
            summary["TXL < 10"].sum(),

        "10 ≤ TXL ≤ 20":
            summary["10 ≤ TXL ≤ 20"].sum(),

        "20 < TXL ≤ 40":
            summary["20 < TXL ≤ 40"].sum(),

        "Đã quá hạn":
            summary["Đã quá hạn"].sum(),

        "Tổng":
            summary["Tổng"].sum()

    }

    summary = pd.concat(
        [
            summary,
            pd.DataFrame([total_row])
        ],
        ignore_index=True
    )

    return summary


# ==========================================================
# GET SUMMARY TEXT
# ==========================================================

def build_summary_text(df):
    
    #alert_df = get_alert_dataframe_bh(df)
    total = len(df)
    #alert_total = len(alert_df)
    
    txl10 = (
        df["sla_bh"]
        .eq("TXL < 10")
        .sum()
    )

    txl20 = (
        df["sla_bh"]
        .eq("10 ≤ TXL ≤ 20")
        .sum()
    )

    txl40 = (
        df["sla_bh"]
        .eq("20 < TXL ≤ 40")
        .sum()
    )

    quahan = (
        df["sla_bh"]
        .eq("Đã quá hạn")
        .sum()
    )

    return (
        f"Tổng phiếu: {total}    |    "
        f"TXL < 10: {txl10}    |    "
        f"10 ≤ TXL ≤ 20: {txl20}    |    "
        f"20 < TXL ≤ 40: {txl40}    |    "
        f"Đã quá hạn: {quahan}"
    )

# ==========================================================
# PROCESS BAO HONG FILE
# ==========================================================

def process_bh_file(filepath):

    try:

        df = pd.read_excel(
            filepath,
            dtype=object
        )

    except Exception as e:

        raise Exception(
            f"Không đọc được file:\n{e}"
        )

    if len(df) == 0:

        df["thoigian_xl_bh"] = []
        df["sla_bh"] = []
        df["time_factory_bh"] = []
        df["time_txl_bh"] = []
        #df["DONVI"] = []

        return df

    # ======================================================
    # Chuẩn hóa dữ liệu
    # ======================================================

    df.columns = [
        str(col).strip()
        for col in df.columns
    ]

    # ======================================================
    # Service --> Chuyển sang giám sát ALL services
    # ======================================================

    #service_mask = (
    #    df["loaihinh_tb"]
    #    .fillna("")
    #    .astype(str)
    #    .apply(is_service_match)
    #)

    # ======================================================
    # Trạng thái
    # ======================================================

    status_mask = (

        #df["ten_trangthai"]
        df["trangthai_bh"]
        .fillna("")
        .astype(str)
        .str.strip()
        .eq("Mới tiếp nhận")

    )

    # ======================================================
    # Đơn vị xử lý rỗng
    # ======================================================

    dvxl_mask = (

        df["ten_dv_xl"]
        .fillna("")
        .astype(str)
        .str.strip()
        .eq("")

    )

    # ======================================================
    # Đơn vị đang thực hiện
    # ======================================================

    donvi_mask = (

        df["ten_dv_dang_th"]
        .fillna("")
        .astype(str)
        .apply(has_donvi_keyword)

    )

    # ======================================================
    # Lọc dữ liệu
    # ======================================================

    df = df[
        #service_mask
        status_mask
        & dvxl_mask
        & donvi_mask
    ].copy()

    if len(df) == 0:

        return df

    # ======================================================
    # Parse ngày báo hỏng
    # ======================================================

    df["ngay_bh_dt"] = (
        df["ngay_bh"]
        .apply(parse_datetime)
    )

    now = datetime.now()

    # ======================================================
    # Thời gian xử lý
    # ======================================================

    df["thoigian_xl_bh"] = (

        (
            now
            - df["ngay_bh_dt"]
        ).dt.total_seconds()

        / 60

    )

    df["thoigian_xl_bh"] = (

        df["thoigian_xl_bh"]
        .fillna(0)
        .round(0)
        .astype(int)

    )

    # ======================================================
    # SLA
    # ======================================================

    df["sla_bh"] = (

        df["thoigian_xl_bh"]
        .apply(calculate_sla_bh)

    )

    # ======================================================
    # Factory Time
    # ======================================================

    df["time_factory_bh"] = (

        df["ngay_bh_dt"]
        + timedelta(minutes=20)

    )

    df["time_factory_bh"] = (

        df["time_factory_bh"]
        .apply(format_datetime)

    )

    # ======================================================
    # TXL Time
    # ======================================================

    df["time_txl_bh"] = (

        df["ngay_bh_dt"]
        + timedelta(minutes=40)

    )

    df["time_txl_bh"] = (

        df["time_txl_bh"]
        .apply(format_datetime)

    )

    # ======================================================
    # Map đơn vị
    # ======================================================

    df["DONVI"] = (

        df["ten_dv"]
        .fillna("")
        .astype(str)
        .apply(map_donvi)

    )

    # ======================================================
    # Sắp xếp ưu tiên
    # Quá hạn -> 20-40 -> 10-20
    # ======================================================

    sla_order = {

        "Đã quá hạn": 1,

        "20 < TXL ≤ 40": 2,

        "10 ≤ TXL ≤ 20": 3,
        
        "TXL < 10":4

    }

    df["sort_sla"] = (

        df["sla_bh"]
        .map(sla_order)

    )

    df = df.sort_values(

        by=[
            "sort_sla",
            "thoigian_xl_bh"
        ],

        ascending=[
            True,
            False
        ]

    )

    df = df.reset_index(
        drop=True
    )

    return df

# ==========================================================
# PROCESS HO TRO FILE
# ==========================================================

def process_ht_file(filepath):

    df = pd.read_excel(
        filepath,
        dtype=object
    )

    df.columns = [
        str(c).strip()
        for c in df.columns
    ]

    #service_mask = (
    #    df["LOAIHINH_TB"]
    #    .fillna("")
    #    .astype(str)
    #    .apply(is_service_match)
    #)

    status_mask = (
        df["TRANGTHAI_HT"]
        .fillna("")
        .astype(str)
        .str.strip()
        .eq("Mới tiếp nhận")
    )

    dvxl_mask = (
        df["DONVI_XULY"]
        .fillna("")
        .astype(str)
        .str.strip()
        .eq("")
    )

    df = df[
        #service_mask
        status_mask
        & dvxl_mask
    ].copy()

    if len(df) == 0:

        df["thoigian_xl_ht"] = []
        df["sla_ht"] = []
        df["time_factory_ht"] = []
        df["time_txl_ht"] = []
        #df["DONVI"] = []

        return df

    df["NGAY_TN_DT"] = (
        df["NGAY_TN"]
        .apply(parse_datetime)
    )

    now = datetime.now()

    df["thoigian_xl_ht"] = (
        (
            now
            - df["NGAY_TN_DT"]
        ).dt.total_seconds()
        / 60
    )

    df["thoigian_xl_ht"] = (
        df["thoigian_xl_ht"]
        .fillna(0)
        .round(0)
        .astype(int)
    )

    df["sla_ht"] = (
        df["thoigian_xl_ht"]
        .apply(calculate_sla_bh)
    )

    df["time_factory_ht"] = (
        df["NGAY_TN_DT"]
        + timedelta(minutes=20)
    )

    df["time_factory_ht"] = (
        df["time_factory_ht"]
        .apply(format_datetime)
    )

    df["time_txl_ht"] = (
        df["NGAY_TN_DT"]
        + timedelta(minutes=40)
    )

    df["time_txl_ht"] = (
        df["time_txl_ht"]
        .apply(format_datetime)
    )

    df["DONVI"] = (
        df["TEN_DV"]
        .fillna("")
        .astype(str)
        .apply(map_donvi)
    )

    sla_order = {

        "Đã quá hạn": 1,
        "20 < TXL ≤ 40": 2,
        "10 ≤ TXL ≤ 20": 3,
        "TXL < 10":4
    }

    df["sort_sla"] = (
        df["sla_ht"]
        .map(sla_order)
    )

    df = df.sort_values(
        by=[
            "sort_sla",
            "thoigian_xl_ht"
        ],
        ascending=[
            True,
            False
        ]
    )

    return df.reset_index(
        drop=True
    )
    
# ==========================================================
# GET ALERT DATAFRAME
# ==========================================================
    
def get_alert_dataframe_bh(df):

    if len(df) == 0:
        return df.copy()

    return df[
        df["thoigian_xl_bh"] >= 10
    ].copy()


def get_alert_dataframe_ht(df):

    if len(df) == 0:
        return df.copy()

    return df[
        df["thoigian_xl_ht"] >= 10
    ].copy()


def get_operational_alert_data(filepath):
    """Return the completion and in-progress alerts from one OneBSS export."""
    try:
        df = pd.read_excel(filepath, dtype=object)
    except Exception as exc:
        raise RuntimeError(f"Không đọc được file cảnh báo vận hành: {exc}") from exc

    df.columns = [str(column).strip() for column in df.columns]
    required = {"ma_bh", "loaihinh_tb", "ten_nv", "ten_dv", "ngay_bh", "trangthai_bh", "ten_dv_xl"}
    missing = sorted(required - set(df.columns))
    if missing:
        raise RuntimeError("File Excel thiếu cột: " + ", ".join(missing))
    if df.empty:
        return df.copy(), df.copy()

    df = df.copy()
    df["ngay_bh_dt"] = df["ngay_bh"].apply(parse_datetime)
    df["thoigian_xl_bh"] = ((datetime.now() - df["ngay_bh_dt"]).dt.total_seconds() / 60).fillna(0).round(0).astype(int)
    df["DONVI"] = df["ten_dv"].fillna("").astype(str).apply(map_donvi)
    status = df["trangthai_bh"].fillna("").astype(str).str.strip()
    handler = df["ten_dv_xl"].fillna("").astype(str).str.strip()
    completion = df[status.isin({"Đã xử lý xong", "VNPT-NET xử lý xong"})].copy()
    completion["alert_status"] = completion["trangthai_bh"].astype(str).str.strip()
    it_completion = df[(status == "Mới tiếp nhận") & (handler == "VNPT IT")].copy()
    it_completion["alert_status"] = "VNPT IT đã xử lý xong"
    completion = pd.concat([completion, it_completion], ignore_index=True)
    in_progress = df[(status.isin({"VNPT-NET xử lý mạng VT lớp trên", "Đang chờ xử lý mạng VT lớp trên", "Đang điều hành xử lý"})) & (df["thoigian_xl_bh"] >= 120)].copy()
    in_progress["alert_status"] = "Đang thực hiện"
    in_progress["alert_round"] = ((in_progress["thoigian_xl_bh"] - 120) // 60 + 1).astype(int)
    return completion, in_progress


def _build_operational_message(title, df, recommendation, include_round=False):
    lines = [f"<b>🔔 {title}</b>", f"<i>{build_timestamp_text()}</i>", ""]
    for index, (_, row) in enumerate(df.iterrows(), start=1):
        value = lambda key: html.escape(str(row.get(key, "") or ""))
        lines.extend([f"{index}. {value('ma_bh')} | {value('loaihinh_tb')}", f"   Người mở phiếu: {value('ten_nv')}", f"   Đơn vị: {value('DONVI')}", f"   Ngày mở phiếu: {value('ngay_bh')}", f"   Trạng thái phiếu: {value('alert_status')}"])
        if include_round:
            lines.append(f"   Cảnh báo lần {int(row['alert_round'])}: đã xử lý {int(row['thoigian_xl_bh'])} phút")
        lines.extend(["   📌 Đề nghị:", f"   • {recommendation}", ""])
    lines.append(f"<b>Tổng số phiếu:</b> {len(df)}")
    return "\n".join(lines)


def build_completion_message(df):
    return _build_operational_message("DANH SÁCH PHIẾU CHƯA NGHIỆM THU", df, "Khẩn trương hoàn công phiếu")


def build_in_progress_message(df):
    return _build_operational_message("DANH SÁCH PHIẾU ĐANG THỰC HIỆN", df, "Khẩn trương hoàn tất xử lý", include_round=True)


def get_under10_dataframe_bh(df):

    if len(df) == 0:
        return df.copy()

    return df[
        df["thoigian_xl_bh"] < 10
    ].copy()


def get_under10_dataframe_ht(df):

    if len(df) == 0:
        return df.copy()

    return df[
        df["thoigian_xl_ht"] < 10
    ].copy()

# ==========================================================
# GET DISPLAY DATAFRAME
# ==========================================================

def get_display_dataframe(df):

    if len(df) == 0:

        return pd.DataFrame()

    display_cols = [

        col_name

        for _, col_name

        in DISPLAY_COLUMNS_BH

    ]

    return df[
        display_cols
    ].copy()


# ==========================================================
# GET UNIT DATAFRAME
# ==========================================================

def get_unit_dataframe(
    df,
    unit_name
):

    if len(df) == 0:

        return df.copy()

    return df[
        df["DONVI"] == unit_name
    ].copy()


# ==========================================================
# GET EXPORT DATAFRAME
# ==========================================================

def get_export_dataframe(df, display_columns):

    display_cols = [

        field

        for _, field

        in display_columns

    ]

    return df.reindex(
        columns=display_cols
    ).copy()

# ==========================================================
# BUILD EXPORT TITLE
# ==========================================================

def build_sheet_title(unit_name=None):

    if unit_name is None:
        return "THỐNG KÊ SLA PHIẾU TXL BÁO HỎNG"

    return f"PHIẾU TXL BÁO HỎNG CỦA {unit_name}"

# ==========================================================
# BUILD TIMESTAMP TEXT
# ==========================================================

def build_timestamp_text():

    global FILE_DATA_TIME

    dt = FILE_DATA_TIME

    if dt is None:
        dt = datetime.now()

    return (

        "(Số liệu cập nhật lúc "

        + dt.strftime(
            "%H:%M ngày %d/%m/%Y"
        )

        + ")"

    )
    
# ==========================================================
# EXPORT EXCEL BH
# ==========================================================

def export_bh_excel(
    df,
    save_path
):

    # ============================================
    # Tách dữ liệu
    # ============================================

    alert_df = get_alert_dataframe_bh(df)

    under10_df = get_under10_dataframe_bh(df)

    if len(alert_df) == 0 and len(under10_df) == 0:

        raise Exception(
            "Không có dữ liệu để export"
        )

    wb = Workbook()

    default_sheet = wb.active

    wb.remove(default_sheet)

    # ======================================================
    # SHEET TONG HOP
    # ======================================================

    summary_df = build_summary_dataframe(alert_df)

    ws = wb.create_sheet(
        "Tonghop"
    )

    # ======================================================
    # ROW 1
    # ======================================================

    ws.merge_cells(
        "A1:E1"
    )

    ws["A1"] = (
        build_sheet_title()
    )

    ws["A1"].font = Font(
        bold=True,
        size=14
    )

    ws["A1"].alignment = Alignment(
        horizontal="center"
    )

    # ======================================================
    # ROW 2
    # ======================================================

    ws.merge_cells(
        "A2:E2"
    )

    ws["A2"] = (
        build_timestamp_text()
    )

    ws["A2"].font = Font(
        italic=True,
        color="FF0000"
    )

    ws["A2"].alignment = Alignment(
        horizontal="center"
    )

    # ======================================================
    # ROW 3 HEADER
    # ======================================================

    headers = [

        "Đơn vị",

        "10 ≤ TXL ≤ 20",

        "20 < TXL ≤ 40",

        "Đã quá hạn",

        "Tổng"

    ]

    for col_idx, header in enumerate(
        headers,
        start=1
    ):

        ws.cell(
            3,
            col_idx,
            header
        )

    style_header_row(
        ws,
        3
    )

    # ======================================================
    # DATA
    # ======================================================

    excel_row = 4

    for _, row in summary_df.iterrows():

        ws.cell(
            excel_row,
            1,
            row["Đơn vị"]
        )

        ws.cell(
            excel_row,
            2,
            row["10 ≤ TXL ≤ 20"]
        )

        ws.cell(
            excel_row,
            3,
            row["20 < TXL ≤ 40"]
        )

        ws.cell(
            excel_row,
            4,
            row["Đã quá hạn"]
        )

        ws.cell(
            excel_row,
            5,
            row["Tổng"]
        )

        # Dòng tổng
        if row["Đơn vị"] == "TỔNG":

            for cell in ws[excel_row]:

                cell.fill = TOTAL_FILL

                cell.font = Font(
                    bold=True
                )

        excel_row += 1

    autosize_columns(ws)

    # ======================================================
    # 3 SHEET DON VI
    # ======================================================

    for donvi in [

        "P.HTKH MB",

        "P.HTKH MT",

        "P.HTKH MN"

    ]:

        unit_df = get_unit_dataframe(
            alert_df,
            donvi
        )

        ws = wb.create_sheet(
            donvi
        )

        # ==============================================
        # ROW 1
        # ==============================================

        total_cols = len(
            DISPLAY_COLUMNS_BH
        )

        ws.merge_cells(

            start_row=1,
            start_column=1,

            end_row=1,
            end_column=total_cols

        )

        ws["A1"] = (
            build_sheet_title(
                donvi
            )
        )

        ws["A1"].font = Font(
            bold=True,
            size=14
        )

        ws["A1"].alignment = Alignment(
            horizontal="center"
        )

        # ==============================================
        # ROW 2
        # ==============================================

        ws.merge_cells(

            start_row=2,
            start_column=1,

            end_row=2,
            end_column=total_cols

        )

        ws["A2"] = (
            build_timestamp_text()
        )

        ws["A2"].font = Font(
            italic=True,
            color="FF0000"
        )

        ws["A2"].alignment = Alignment(
            horizontal="center"
        )

        # ==============================================
        # ROW 3 HEADER
        # ==============================================

        for col_idx, (

            header,

            field

        ) in enumerate(
            DISPLAY_COLUMNS_BH,
            start=1
        ):

            ws.cell(
                3,
                col_idx,
                header
            )

        style_header_row(
            ws,
            3
        )

        # ==============================================
        # ROW 4+
        # ==============================================

        export_df = get_export_dataframe(
            unit_df,
            DISPLAY_COLUMNS_BH
        )

        current_row = 4

        for _, row in export_df.iterrows():

            for col_idx, (

                header,

                field

            ) in enumerate(
                DISPLAY_COLUMNS_BH,
                start=1
            ):

                value = row.get(
                    field,
                    ""
                )

                ws.cell(
                    current_row,
                    col_idx,
                    value
                )

            # Highlight SLA

            apply_sla_fill(

                ws,

                current_row,

                row["sla_bh"]

            )

            current_row += 1

        autosize_columns(ws)
        
    # ======================================================
    # SHEET TXL < 10
    # ======================================================

    ws = wb.create_sheet("TXL < 10")

    total_cols = len(DISPLAY_COLUMNS_BH)

    # ------------------------------------------------------
    # ROW 1
    # ------------------------------------------------------

    ws.merge_cells(

        start_row=1,
        start_column=1,

        end_row=1,
        end_column=total_cols

    )

    ws["A1"] = "PHIẾU TXL BÁO HỎNG < 10 PHÚT"

    ws["A1"].font = Font(
        bold=True,
        size=14
    )

    ws["A1"].alignment = Alignment(
        horizontal="center"
    )

    # ------------------------------------------------------
    # ROW 2
    # ------------------------------------------------------

    ws.merge_cells(

        start_row=2,
        start_column=1,

        end_row=2,
        end_column=total_cols

    )

    ws["A2"] = build_timestamp_text()

    ws["A2"].font = Font(
        italic=True,
        color="FF0000"
    )

    ws["A2"].alignment = Alignment(
        horizontal="center"
    )

    # ------------------------------------------------------
    # ROW 3 HEADER
    # ------------------------------------------------------

    for col_idx, (

        header,

        field

    ) in enumerate(

        DISPLAY_COLUMNS_BH,

        start=1

    ):

        ws.cell(
            3,
            col_idx,
            header
        )

    style_header_row(
        ws,
        3
    )

    # ------------------------------------------------------
    # ROW 4+
    # ------------------------------------------------------

    export_df = get_export_dataframe(
        under10_df,
        DISPLAY_COLUMNS_BH
    )

    current_row = 4

    for _, row in export_df.iterrows():

        for col_idx, (

            header,

            field

        ) in enumerate(

            DISPLAY_COLUMNS_BH,

            start=1

        ):

            value = row.get(
                field,
                ""
            )

            ws.cell(
                current_row,
                col_idx,
                value
            )

        current_row += 1

    autosize_columns(ws)

    # ======================================================
    # SAVE
    # ======================================================

    wb.save(
        save_path
    )
    
def export_ht_excel(df, save_path):

    alert_df = get_alert_dataframe_ht(df)

    under10_df = get_under10_dataframe_ht(df)

    if len(alert_df) == 0 and len(under10_df) == 0:

        raise Exception(
            "Không có dữ liệu để export"
        )

    wb = Workbook()

    wb.remove(wb.active)

    summary_df = build_summary_dataframe_ht(alert_df)

    # ==================================================
    # SHEET TONG HOP
    # ==================================================

    ws = wb.create_sheet("Tonghop")

    ws.merge_cells("A1:E1")

    ws["A1"] = "THỐNG KÊ SLA PHIẾU TXL HỖ TRỢ"

    ws["A1"].font = Font(
        bold=True,
        size=14
    )

    ws["A1"].alignment = Alignment(
        horizontal="center"
    )

    ws.merge_cells("A2:E2")

    ws["A2"] = build_timestamp_text()

    ws["A2"].font = Font(
        italic=True,
        color="FF0000"
    )

    ws["A2"].alignment = Alignment(
        horizontal="center"
    )

    headers = [

        "Đơn vị",
        "10 ≤ TXL ≤ 20",
        "20 < TXL ≤ 40",
        "Đã quá hạn",
        "Tổng"

    ]

    for col_idx, header in enumerate(
        headers,
        start=1
    ):
        ws.cell(
            3,
            col_idx,
            header
        )

    style_header_row(ws, 3)

    excel_row = 4

    for _, row in summary_df.iterrows():

        ws.cell(excel_row,1,row["Đơn vị"])
        ws.cell(excel_row,2,row["10 ≤ TXL ≤ 20"])
        ws.cell(excel_row,3,row["20 < TXL ≤ 40"])
        ws.cell(excel_row,4,row["Đã quá hạn"])
        ws.cell(excel_row,5,row["Tổng"])

        if row["Đơn vị"] == "TỔNG":

            for cell in ws[excel_row]:

                cell.fill = TOTAL_FILL
                cell.font = Font(bold=True)

        excel_row += 1

    autosize_columns(ws)

    # ==================================================
    # 3 SHEET DON VI
    # ==================================================

    for donvi in [

        "P.HTKH MB",
        "P.HTKH MT",
        "P.HTKH MN"

    ]:

        unit_df = get_unit_dataframe(
            alert_df,
            donvi
        )

        ws = wb.create_sheet(donvi)

        total_cols = len(
            DISPLAY_COLUMNS_HT
        )

        ws.merge_cells(
            start_row=1,
            start_column=1,
            end_row=1,
            end_column=total_cols
        )

        ws["A1"] = (
            f"PHIẾU TXL HỖ TRỢ CỦA {donvi}"
        )

        ws["A1"].font = Font(
            bold=True,
            size=14
        )

        ws["A1"].alignment = Alignment(
            horizontal="center"
        )

        ws.merge_cells(
            start_row=2,
            start_column=1,
            end_row=2,
            end_column=total_cols
        )

        ws["A2"] = build_timestamp_text()

        ws["A2"].font = Font(
            italic=True,
            color="FF0000"
        )

        ws["A2"].alignment = Alignment(
            horizontal="center"
        )

        for col_idx, (
            header,
            field
        ) in enumerate(
            DISPLAY_COLUMNS_HT,
            start=1
        ):

            ws.cell(
                3,
                col_idx,
                header
            )

        style_header_row(ws,3)
        
        export_df = get_export_dataframe(
            unit_df,
            DISPLAY_COLUMNS_HT
        )

        current_row = 4

        for _, row in export_df.iterrows():

            for col_idx, (
                header,
                field
            ) in enumerate(
                DISPLAY_COLUMNS_HT,
                start=1
            ):

                ws.cell(
                    current_row,
                    col_idx,
                    row.get(field,"")
                )

            apply_sla_fill(
                ws,
                current_row,
                row["sla_ht"]
            )

            current_row += 1

        autosize_columns(ws)
        
    # ======================================================
    # SHEET TXL < 10
    # ======================================================

    ws = wb.create_sheet("TXL < 10")

    total_cols = len(DISPLAY_COLUMNS_HT)

    # ------------------------------------------------------
    # ROW 1
    # ------------------------------------------------------

    ws.merge_cells(

        start_row=1,
        start_column=1,

        end_row=1,
        end_column=total_cols

    )

    ws["A1"] = "PHIẾU TXL HỖ TRỢ < 10 PHÚT"

    ws["A1"].font = Font(
        bold=True,
        size=14
    )

    ws["A1"].alignment = Alignment(
        horizontal="center"
    )

    # ------------------------------------------------------
    # ROW 2
    # ------------------------------------------------------

    ws.merge_cells(

        start_row=2,
        start_column=1,

        end_row=2,
        end_column=total_cols

    )

    ws["A2"] = build_timestamp_text()

    ws["A2"].font = Font(
        italic=True,
        color="FF0000"
    )

    ws["A2"].alignment = Alignment(
        horizontal="center"
    )

    # ------------------------------------------------------
    # ROW 3 HEADER
    # ------------------------------------------------------

    for col_idx, (

        header,

        field

    ) in enumerate(

        DISPLAY_COLUMNS_HT,

        start=1

    ):

        ws.cell(
            3,
            col_idx,
            header
        )

    style_header_row(
        ws,
        3
    )

    # ------------------------------------------------------
    # ROW 4+
    # ------------------------------------------------------

    export_df = get_export_dataframe(
        under10_df,
        DISPLAY_COLUMNS_HT
    )

    current_row = 4

    for _, row in export_df.iterrows():

        for col_idx, (

            header,

            field

        ) in enumerate(

            DISPLAY_COLUMNS_HT,

            start=1

        ):

            value = row.get(
                field,
                ""
            )

            ws.cell(
                current_row,
                col_idx,
                value
            )

        current_row += 1

    autosize_columns(ws)

    wb.save(save_path)

# ==========================================================
# Hàm gửi Telegram
# ==========================================================
    
def send_telegram_message(text, *, chat_id=None, bot_token=None):

    # Read the environment again at send time. This also supports a long-lived
    # app process whose Telegram configuration is refreshed before a run. An
    # explicit chat ID is used for private workflow-error notifications.
    bot_token = (bot_token or os.getenv("TXL_TELEGRAM_BOT_TOKEN", BOT_TOKEN)).strip()
    target_chat_id = str(
        chat_id if chat_id is not None
        else os.getenv("TXL_TELEGRAM_GROUP_CHAT_ID", GROUP_CHAT_ID)
    ).strip()

    if not bot_token or not target_chat_id:
        raise RuntimeError(
            "Chưa cấu hình Telegram Bot token hoặc chat ID đích"
        )

    url = (
        f"https://api.telegram.org/bot"
        f"{bot_token}/sendMessage"
    )

    payload = {

        "chat_id": target_chat_id,

        "text": text,
        
        "parse_mode": "HTML"

    }

    response = requests.post(
        url,
        data=payload,
        timeout=15
    )

    try:
        result = response.json()
    except ValueError:
        result = {}

    if not response.ok or not result.get("ok"):
        # Never include the request URL here because it contains the bot token.
        description = result.get("description") or f"HTTP {response.status_code}"
        raise RuntimeError(f"Gửi Telegram thất bại: {description}")

    return True
    
# ==========================================================
# BUILD TELEGRAM MESSAGE BH
# ==========================================================

def build_bh_message(df):

    lines = []
    
    count = 0

    lines.append(
        "<b>🔔 DANH SÁCH PHIẾU TIỀN XỬ LÝ BÁO HỎNG</b>"
    )

    lines.append(
        #build_timestamp_text()
        f"<i>{build_timestamp_text()}</i>"
    )

    lines.append("")

    for idx, (_, row) in enumerate(
        df.iterrows(),
        start=1
    ):

        sla = row["sla_bh"]

        if sla == "10 ≤ TXL ≤ 20":
            
            count += 1
            
            lines.append(
                f"{idx}. 🟢 {row['ma_bh']} | {row['loaihinh_tb']}"
            )

            lines.append(
                f"   Người mở phiếu: {row['ten_nv']}"
            )
            
            lines.append(
                f"   Đơn vị: {row['DONVI']}"
            )
            
            lines.append(
                f"   Ngày mở phiếu: "
                f"{row['ngay_bh']}"
            )

            lines.append(
                f"   Thời gian đã xử lý: "
                f"{row['thoigian_xl_bh']} phút"
            )
            
            lines.append(
                f"   📌 Đề nghị:"
            )

            lines.append(
                f"   • Chuyển Factory trước: "
                f"{row['time_factory_bh']} (TXLBE ≤ 20 phút)."
            )

            lines.append(
                f"   • Hoặc khóa phiếu trước: "
                f"{row['time_txl_bh']} (TXLTC ≤ 40 phút)."
            )

        elif sla == "20 < TXL ≤ 40":
            
            count += 1

            lines.append(
                f"{idx}. 🟡 {row['ma_bh']} | {row['loaihinh_tb']}"
            )

            lines.append(
                f"   Người mở phiếu: {row['ten_nv']}"
            )
            
            lines.append(
                f"   Đơn vị: {row['DONVI']}"
            )
            
            lines.append(
                f"   Ngày mở phiếu: "
                f"{row['ngay_bh']}"
            )

            lines.append(
                f"   Thời gian đã xử lý: "
                f"{row['thoigian_xl_bh']} phút/40 phút"
            )
            
            lines.append(
                f"   📌 Đề nghị:"
            )
            
            lines.append(
                f"   • Khẩn trương hoàn tất xử lý"
            )

            lines.append(
                f"   • Khóa phiếu trước  "
                f"{row['time_txl_bh']} để đảm bảo SLA (TXLTC ≤ 40 phút)."
            )

        else:
            
            count += 1

            lines.append(
                f"{idx}. 🔴 {row['ma_bh']} | {row['loaihinh_tb']}"
            )

            lines.append(
                f"   Người mở phiếu: {row['ten_nv']}"
            )
            
            lines.append(
                f"   Đơn vị: {row['DONVI']}"
            )
            
            lines.append(
                f"   Ngày mở phiếu: "
                f"{row['ngay_bh']}"
            )

            lines.append(
                f"   Thời gian đã xử lý: "
                f"{row['thoigian_xl_bh']} phút/40 phút"
            )
            
            lines.append(
                f"   🆘 Đã quá hạn thời gian tiền xử lý (TXLTC ≤ 40 phút)."
            )            
            
            lines.append(
                f"   📌 Đề nghị:"
            )

            lines.append(
                "   • Kiểm tra trạng thái xử lý"
            )
            
            lines.append(
                "   • Cập nhật tiến trình, bổ sung nguyên nhân kéo dài"
            )          

        lines.append("")
        
    if count == 0:
        return build_no_ticket_message("BH")

    #lines.append("")
    lines.append(
        f"<b>Tổng số phiếu:</b> {count}"
    )

    return "\n".join(lines)

# ==========================================================
# BUILD TELEGRAM MESSAGE HT
# ==========================================================

def build_ht_message(df):

    lines = []
    
    count = 0

    lines.append(
        "<b>🔔 DANH SÁCH PHIẾU TIỀN XỬ LÝ HỖ TRỢ</b>"
    )

    lines.append(
        #build_timestamp_text()
        f"<i>{build_timestamp_text()}</i>"
    )

    lines.append("")

    for idx, (_, row) in enumerate(
        df.iterrows(),
        start=1
    ):

        sla = row["sla_ht"]

        if sla == "10 ≤ TXL ≤ 20":
                
            count += 1

            lines.append(
                f"{idx}. 🟢 {row['MA_HT']} | {row['LOAIHINH_TB']}"
            )

            lines.append(
                f"   Người mở phiếu: {row['TEN_NV']}"
            )

            lines.append(
                f"   Đơn vị: {row['DONVI']}"
            )
            
            lines.append(
                f"   Ngày tiếp nhận: "
                f"{row['NGAY_TN']}"
            )

            lines.append(
                f"   Thời gian đã xử lý: "
                f"{row['thoigian_xl_ht']} phút"
            )
            
            lines.append(
                f"   📌 Đề nghị:"
            )            

            lines.append(
                f"   • Chuyển Factory trước "
                f"{row['time_factory_ht']} (TXLBE ≤ 20 phút)."
            )

            lines.append(
                f"   • Hoặc khóa phiếu trước "
                f"{row['time_txl_ht']} (TXLTC ≤ 40 phút)."
            )

        elif sla == "20 < TXL ≤ 40":
            
            count += 1

            lines.append(
                f"{idx}. 🟡 {row['MA_HT']} | {row['LOAIHINH_TB']}"
            )

            lines.append(
                f"   Người mở phiếu: {row['TEN_NV']}"
            )

            lines.append(
                f"   Đơn vị: {row['DONVI']}"
            )
            
            lines.append(
                f"   Ngày tiếp nhận: "
                f"{row['NGAY_TN']}"
            )
            
            lines.append(
                f"   Thời gian đã xử lý: "
                f"{row['thoigian_xl_ht']} phút/40 phút"
            )
            
            lines.append(
                f"   📌 Đề nghị:"
            )
            
            lines.append(
                f"   • Khẩn trương hoàn tất xử lý"
            )            

            lines.append(
                f"   • Khóa phiếu trước "
                f"{row['time_txl_ht']} để đảm bảo SLA (TXLTC ≤ 40 phút)."
            )

        else:
            
            count += 1

            lines.append(
                f"{idx}. 🔴 {row['MA_HT']} | {row['LOAIHINH_TB']}"
            )

            lines.append(
                f"   Người mở phiếu: {row['TEN_NV']}"
            )

            lines.append(
                f"   Đơn vị: {row['DONVI']}"
            )
            
            lines.append(
                f"   Ngày tiếp nhận: "
                f"{row['NGAY_TN']}"
            )

            lines.append(
                f"   Thời gian đã xử lý: "
                f"{row['thoigian_xl_ht']} phút/40 phút"
            )

            lines.append(
                f"   🆘 Đã quá hạn thời gian tiền xử lý (TXLTC ≤ 40 phút)."
            )            
            
            lines.append(
                f"   📌 Đề nghị:"
            )

            lines.append(
                "   • Kiểm tra trạng thái xử lý"
            )
            
            lines.append(
                "   • Cập nhật tiến trình, bổ sung nguyên nhân kéo dài"
            )

        lines.append("")
        
    if count == 0:
        return build_no_ticket_message("HT")

    lines.append(
        f"<b>Tổng số phiếu:</b> {count}"
    )

    return "\n".join(lines)
    
# ==========================================================
# BUILD TELEGRAM MESSAGE KHI KHÔNG CÓ DỮ LIỆU TXL
# ==========================================================

def build_no_ticket_message(data_type):

    if data_type == "BH":

        return "\n".join([

            "<b>🔔 GIÁM SÁT PHIẾU TXL BÁO HỎNG</b>",

            #build_timestamp_text(),
            f"<i>{build_timestamp_text()}</i>",

            "",

            "Hiện không có phiếu Tiền xử lý báo hỏng nào có thời gian xử lý ≥ 10 phút vào thời điểm này."

        ])

    else:

        return "\n".join([

            "<b>🔔 GIÁM SÁT PHIẾU TXL HỖ TRỢ</b>",

            #build_timestamp_text(),
            f"<i>{build_timestamp_text()}</i>",

            "",

            "Hiện không có phiếu Tiền xử lý hỗ trợ nào có thời gian xử lý ≥ 10 phút vào thời điểm này."

        ])
    
# ==========================================================
# GUI
# ==========================================================

class TXLApp:

    def __init__(self, root):

        self.root = root

        self.root.title(APP_TITLE)

        self.root.geometry("1600x850")

        self.df_bh = pd.DataFrame()
        
        self.df_ht = pd.DataFrame()

        self.data_type = ""

        self.current_file = ""

        self.build_ui()

    # ======================================================
    # UI
    # ======================================================

    def build_ui(self):

        # --------------------------------------------------
        # TOP
        # --------------------------------------------------

        top_frame = tk.Frame(
            self.root
        )

        top_frame.pack(
            fill="x",
            pady=10
        )

        self.btn_upload = tk.Button(

            top_frame,

            text="📂 Chọn File Upload",

            font=(
                "Arial",
                11,
                "bold"
            ),

            bg="#FFE599",              # vàng nhạt

            activebackground="#FFD966",

            relief="raised",

            bd=2,

            padx=10,

            pady=3,

            command=self.select_file

        )

        self.btn_upload.pack()

        # --------------------------------------------------
        # TITLE
        # --------------------------------------------------

        self.lbl_title = tk.Label(

            self.root,

            text="",

            font=(
                "Arial",
                16,
                "bold"
            )

        )

        self.lbl_title.pack(
            pady=(10, 2)
        )

        # --------------------------------------------------
        # TIMESTAMP
        # --------------------------------------------------

        self.lbl_time = tk.Label(

            self.root,

            text="",

            fg="red",

            font=(
                "Arial",
                10,
                "italic"
            )

        )

        self.lbl_time.pack()

        # --------------------------------------------------
        # SUMMARY
        # --------------------------------------------------

        self.lbl_summary = tk.Label(

            self.root,

            text="",

            fg="#333333",

            font=(
                "Arial",
                10
            )

        )

        self.lbl_summary.pack(
            pady=5
        )

        # --------------------------------------------------
        # TABLE FRAME
        # --------------------------------------------------

        table_frame = tk.Frame(
            self.root
        )

        table_frame.pack(

            fill="both",

            expand=True,

            padx=10,

            pady=10

        )

        self.tree = ttk.Treeview(

            table_frame,

            show="headings",

            selectmode="extended"

        )

        # --------------------------------------------------
        # HEADERS
        # --------------------------------------------------
        
        style = ttk.Style()

        style.configure(
            "Treeview.Heading",
            font=("Arial", 10, "bold")
        )
        
        # --------------------------------------------------
        # SCROLLBAR
        # --------------------------------------------------

        y_scroll = ttk.Scrollbar(

            table_frame,

            orient="vertical",

            command=self.tree.yview

        )

        x_scroll = ttk.Scrollbar(

            table_frame,

            orient="horizontal",

            command=self.tree.xview

        )

        self.tree.configure(

            yscrollcommand=y_scroll.set,

            xscrollcommand=x_scroll.set

        )

        self.tree.grid(
            row=0,
            column=0,
            sticky="nsew"
        )

        y_scroll.grid(
            row=0,
            column=1,
            sticky="ns"
        )

        x_scroll.grid(
            row=1,
            column=0,
            sticky="ew"
        )

        table_frame.rowconfigure(
            0,
            weight=1
        )

        table_frame.columnconfigure(
            0,
            weight=1
        )

        # --------------------------------------------------
        # TAGS
        # --------------------------------------------------

        self.tree.tag_configure(

            "yellow",

            background="#FFF2CC"

        )

        self.tree.tag_configure(

            "red",

            background="#F4CCCC"

        )

        # --------------------------------------------------
        # BOTTOM
        # --------------------------------------------------

        bottom_frame = tk.Frame(
            self.root
        )

        bottom_frame.pack(
            pady=10
        )

        self.btn_export = tk.Button(

            bottom_frame,

            text="Export phiếu TXL",

            width=20,

            bg="#0094D9",

            fg="white",

            activebackground="#0077B6",

            activeforeground="white",

            font=("Arial", 10, "bold"),

            relief="raised",

            bd=2,

            command=self.export_excel

        )

        self.btn_export.pack(
            side="left",
            padx=10
        )

        self.btn_telegram = tk.Button(

            bottom_frame,

            text="Gởi Telegram",

            width=20,

            bg="#0094D9",

            fg="white",

            activebackground="#0077B6",

            activeforeground="white",

            font=("Arial", 10, "bold"),

            relief="raised",

            bd=2,

            command=self.send_telegram

        )

        self.btn_telegram.pack(
            side="left",
            padx=10
        )
        
    # ======================================================
    # SELECT FILE
    # ======================================================

    def select_file(self):

        filepath = filedialog.askopenfilename(

            filetypes=[
                (
                    "Excel Files",
                    "*.xlsx *.xls"
                )
            ]

        )

        if not filepath:
            return

        self.current_file = filepath
        
        global FILE_DATA_TIME

        FILE_DATA_TIME = get_file_datetime(
            filepath
        )

        # --------------------------------------------------
        # BAO HONG
        # --------------------------------------------------

        if is_valid_bh_file(filepath):

            try:

                self.df_bh = process_bh_file(
                    filepath
                )
                
                self.data_type = "BH"
                
                if len(self.df_bh) == 0:

                    messagebox.showinfo(
                        "Thông báo",
                        "Không có phiếu TXL thỏa điều kiện."
                    )

                    self.tree.delete(
                        *self.tree.get_children()
                    )

                    self.lbl_title.config(
                        text="CHI TIẾT PHIẾU TXL BÁO HỎNG"
                    )

                    self.lbl_time.config(
                        text=build_timestamp_text()
                    )

                    self.lbl_summary.config(
                        text="Tổng phiếu: 0"
                    )

                    return
                
                self.show_bh_data()

                messagebox.showinfo(

                    "Thông báo",

                    f"Tìm thấy "
                    f"{len(self.df_bh)} "
                    f"phiếu TXL"

                )

            except Exception as e:

                messagebox.showerror(
                    "Lỗi",
                    str(e)
                )

        # --------------------------------------------------
        # HO TRO
        # --------------------------------------------------

        elif is_valid_ht_file(filepath):

            try:

                self.df_ht = process_ht_file(
                    filepath
                )

                self.data_type = "HT"
                
                if len(self.df_ht) == 0:

                    messagebox.showinfo(
                        "Thông báo",
                        "Không có phiếu TXL thỏa điều kiện."
                    )

                    self.tree.delete(
                        *self.tree.get_children()
                    )

                    self.lbl_title.config(
                        text="CHI TIẾT PHIẾU TXL HỖ TRỢ"
                    )

                    self.lbl_time.config(
                        text=build_timestamp_text()
                    )

                    self.lbl_summary.config(
                        text="Tổng phiếu: 0"
                    )

                    return

                self.show_ht_data()

                messagebox.showinfo(

                    "Thông báo",

                    f"Tìm thấy "
                    f"{len(self.df_ht)} "
                    f"phiếu TXL"

                )

            except Exception as e:

                messagebox.showerror(

                    "Lỗi",

                    "Tên file không hợp lệ."

                )
    
    # ======================================================
    # CẤU TRÚC TREE VIEW THEO DỮ LIỆU BH/HT
    # ======================================================
    
    def setup_tree_columns(self, display_columns):

        self.tree.delete(*self.tree.get_children())

        self.tree["columns"] = [
            field
            for _, field in display_columns
        ]

        for col in self.tree["columns"]:
            self.tree.heading(col, text="")
            self.tree.column(col, width=0)

        for header, field in display_columns:

            self.tree.heading(
                field,
                text=header
            )

            width = max(
                len(header) * 9,
                140
            )

            self.tree.column(
                field,
                width=width,
                minwidth=80,
                stretch=True
            )

    # ======================================================
    # SHOW DATA
    # ======================================================

    def show_bh_data(self):
        
        self.setup_tree_columns(
            DISPLAY_COLUMNS_BH
        )

        self.tree.delete(
            *self.tree.get_children()
        )

        self.lbl_title.config(

            text=
            "CHI TIẾT PHIẾU TXL BÁO HỎNG"

        )

        self.lbl_time.config(

            text=build_timestamp_text()

        )

        self.lbl_summary.config(

            text=build_summary_text(
                self.df_bh
            )

        )

        display_df = get_display_dataframe(
            self.df_bh
        )

        for _, row in display_df.iterrows():

            values = [

                row.get(
                    field,
                    ""
                )

                for _, field

                in DISPLAY_COLUMNS_BH

            ]

            tag = ""

            sla = row.get(
                "sla_bh",
                ""
            )

            if sla == "20 < TXL ≤ 40":

                tag = "yellow"

            elif sla == "Đã quá hạn":

                tag = "red"

            self.tree.insert(

                "",

                "end",

                values=values,

                tags=(tag,)

            )
            
    def show_ht_data(self):

        self.setup_tree_columns(
            DISPLAY_COLUMNS_HT
        )

        self.tree.delete(
            *self.tree.get_children()
        )

        self.lbl_title.config(
            text="CHI TIẾT PHIẾU TXL HỖ TRỢ"
        )

        self.lbl_time.config(
            text=build_timestamp_text()
        )

        # ==================================================
        # SUMMARY
        # ==================================================

        total = len(self.df_ht)

        txl10 = (
            self.df_ht["sla_ht"]
            .eq("TXL < 10")
            .sum()
        )

        txl20 = (
            self.df_ht["sla_ht"]
            .eq("10 ≤ TXL ≤ 20")
            .sum()
        )

        txl40 = (
            self.df_ht["sla_ht"]
            .eq("20 < TXL ≤ 40")
            .sum()
        )

        qh = (
            self.df_ht["sla_ht"]
            .eq("Đã quá hạn")
            .sum()
        )

        self.lbl_summary.config(

            text=
            f"Tổng phiếu: {total}  |  "
            f"TXL < 10: {txl10}  |  "
            f"10 ≤ TXL ≤ 20: {txl20}  |  "
            f"20 < TXL ≤ 40: {txl40}  |  "
            f"Đã quá hạn: {qh}"

        )

        # ==================================================
        # DETAIL
        # ==================================================

        for _, row in self.df_ht.iterrows():

            values = [

                row.get(field, "")

                for _, field

                in DISPLAY_COLUMNS_HT

            ]

            tag = ""

            if row["sla_ht"] == "20 < TXL ≤ 40":

                tag = "yellow"

            elif row["sla_ht"] == "Đã quá hạn":

                tag = "red"

            self.tree.insert(

                "",

                "end",

                values=values,

                tags=(tag,)

            )

    # ======================================================
    # EXPORT
    # ======================================================

    def export_excel(self):
        
        global FILE_DATA_TIME

        if self.data_type == "BH":

            if len(self.df_bh) == 0:
                return

            dt = FILE_DATA_TIME or datetime.now()

            filename = (
                "Thong_ke_SLA_TXL_BH_"
                + dt.strftime("%d%m%Y_%H%M")
                + ".xlsx"
            )

        elif self.data_type == "HT":

            if len(self.df_ht) == 0:
                return

            dt = FILE_DATA_TIME or datetime.now()

            filename = (
                "Thong_ke_SLA_TXL_HT_"
                + dt.strftime("%d%m%Y_%H%M")
                + ".xlsx"
            )

        else:

            messagebox.showwarning(
                "Thông báo",
                "Chưa có dữ liệu"
            )

            return

        save_path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            initialfile=filename
        )

        if not save_path:
            return

        try:

            if self.data_type == "BH":

                export_bh_excel(
                    self.df_bh,
                    save_path
                )

            else:

                export_ht_excel(
                    self.df_ht,
                    save_path
                )

            messagebox.showinfo(
                "Thành công",
                save_path
            )

        except Exception as e:

            messagebox.showerror(
                "Lỗi",
                str(e)
            )

    # ======================================================
    # SEND TELEGRAM
    # ======================================================

    def send_telegram(self):

        try:

            # ======================================
            # BAO HONG
            # ======================================

            if self.data_type == "BH":

                if len(self.df_bh) == 0:

                    messagebox.showinfo(
                        "Thông báo",
                        "Không có phiếu Tiền xử lý báo hỏng từ 10 phút trở lên. "
                        "Telegram sẽ không được gửi."
                    )

                    return

                alert_df = get_alert_dataframe_bh(self.df_bh)

                if len(alert_df) == 0:

                    messagebox.showinfo(
                        "Thông báo",
                        "Không có phiếu Tiền xử lý báo hỏng từ 10 phút trở lên. "
                        "Telegram sẽ không được gửi."
                    )

                    return

                msg = build_bh_message(alert_df)

                total = len(alert_df)

            # ======================================
            # HO TRO
            # ======================================

            elif self.data_type == "HT":

                if len(self.df_ht) == 0:

                    msg = build_no_ticket_message(
                        "HT"
                    )

                    ok = send_telegram_message(
                        msg
                    )

                    if ok:

                        messagebox.showinfo(
                            "Thông báo",
                            "Đã gửi thông báo không có phiếu TXL"
                        )

                    return
                
                alert_df = get_alert_dataframe_ht(self.df_ht)

                msg = build_ht_message(alert_df)

                total = len(alert_df)

            else:

                messagebox.showwarning(
                    "Thông báo",
                    "Chưa có dữ liệu"
                )

                return

            # ======================================
            # GUI TELEGRAM
            # ======================================

            ok = send_telegram_message(
                msg
            )

            if ok:

                messagebox.showinfo(

                    "Thành công",

                    f"Đã gửi cảnh báo Telegram cho {total} phiếu TXL có thời gian xử lý ≥ 10 phút"

                )

            else:

                messagebox.showerror(

                    "Lỗi",

                    "Không gửi được Telegram"

                )

        except Exception as e:

            messagebox.showerror(

                "Lỗi",

                str(e)

            )

# ==========================================================
# MAIN
# ==========================================================

if __name__ == "__main__":

    root = tk.Tk()

    app = TXLApp(root)

    root.mainloop()
